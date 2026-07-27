"""Girdi cozumleyici: Word (.docx), duz metin/markdown ve JSON -> icerik modeli.

Cozumleyici hicbir sey EKLEMEZ, hicbir sey CIKARMAZ; yalnizca kullanicinin
Word dosyasindaki sirayi ve yapiyi (soru / ornek / konu anlatimi / not,
secenekler, cozum, madde, formul, gorsel) tanir ve modele cevirir.
"""

from __future__ import annotations

import json
import re
import tempfile
from pathlib import Path

from .model import Block, Document, Image, Option, Para, Run

# --------------------------------------------------------------------------
# Isaretler
# --------------------------------------------------------------------------
MARKER_RE = re.compile(r"^\[\s*(?P<name>[^\]]+?)\s*\]\s*", re.IGNORECASE)

_KIND_ALIASES = {
    "soru": "soru",
    "ornek": "ornek",
    "örnek": "ornek",
    "konu": "konu",
    "konu anlatimi": "konu",
    "konu anlatımı": "konu",
    "not": "not",
    "bilgi": "not",
}

_OPTION_RE = re.compile(r"^\(?\s*([A-Ea-e])\s*[).]\s+(.*)$", re.DOTALL)
_SOLUTION_RE = re.compile(r"^(çözüm|cozum)\s*[:.]?\s*", re.IGNORECASE)
_ORNEK_RE = re.compile(r"^(örnek|ornek)\s*(\d+)?\s*[:.\-–]?\s*", re.IGNORECASE)
_SORU_RE = re.compile(r"^soru\s*(\d+)?\s*[:.\-–]?\s*", re.IGNORECASE)
_KONU_RE = re.compile(
    r"^(konu\s*anlat[iı]m[iı]|konu)\s*[:.\-–]\s*", re.IGNORECASE
)
_NOT_RE = re.compile(r"^(not|bilgi\s*notu)\s*[:.\-–]\s*", re.IGNORECASE)
_NUMBERED_RE = re.compile(r"^(\d{1,3})\s*[).]\s+(.*)$", re.DOTALL)
_BULLET_RE = re.compile(r"^[-•*·]\s+(.*)$", re.DOTALL)
_FORMULA_RE = re.compile(r"^\$\$(.*)\$\$$", re.DOTALL)


# --------------------------------------------------------------------------
# Satir ici bicimlendirme
# --------------------------------------------------------------------------
_INLINE_RE = re.compile(
    r"""
    (?P<esc>\\[\\*^_$])
  | (?P<bold>\*\*(?P<bold_text>.+?)\*\*)
  | (?P<supg>\^\{(?P<sup_group>[^}]*)\})
  | (?P<subg>_\{(?P<sub_group>[^}]*)\})
  | (?P<sup>\^(?P<sup_char>[0-9A-Za-zçğıöşüÇĞİÖŞÜ+\-]))
  | (?P<sub>_(?P<sub_char>[0-9A-Za-zçğıöşüÇĞİÖŞÜ+\-]))
    """,
    re.VERBOSE,
)


def parse_inline(text: str, bold: bool = False, script: str = "") -> list[Run]:
    """Duz metindeki hafif isaretlemeyi Run listesine cevirir."""
    runs: list[Run] = []
    pos = 0

    def push(chunk: str, b: bool = bold, s: str = script):
        if chunk:
            runs.append(Run(chunk, bold=b, script=s))

    for m in _INLINE_RE.finditer(text):
        push(text[pos : m.start()])
        if m.group("esc"):
            push(m.group("esc")[1])
        elif m.group("bold"):
            push(m.group("bold_text"), True, script)
        elif m.group("supg"):
            push(m.group("sup_group"), bold, "sup")
        elif m.group("subg"):
            push(m.group("sub_group"), bold, "sub")
        elif m.group("sup"):
            push(m.group("sup_char"), bold, "sup")
        elif m.group("sub"):
            push(m.group("sub_char"), bold, "sub")
        pos = m.end()
    push(text[pos:])
    return runs or [Run("", bold=bold, script=script)]


def _strip_runs(runs: list[Run]) -> list[Run]:
    """Bas/son bosluklari kirpar, bos Run'lari atar."""
    out = [r for r in runs if r.text]
    if out:
        out[0] = Run(out[0].text.lstrip(), out[0].bold, out[0].script)
        out[-1] = Run(out[-1].text.rstrip(), out[-1].bold, out[-1].script)
    return [r for r in out if r.text]


# --------------------------------------------------------------------------
# Ortak cozumleyici cekirdegi
# --------------------------------------------------------------------------
class _Builder:
    """Sirali paragraf akisini bloklara donusturur."""

    def __init__(self):
        self.doc = Document()
        self.block: Block | None = None
        self.in_solution = False
        self.pending_page_break = False
        self.pending_full_page = False

    # -- blok yonetimi --
    def start(self, kind: str, number: int | None = None) -> Block:
        block = Block(kind=kind, number=number)
        block.page_break_before = self.pending_page_break
        block.force_full_page = self.pending_full_page or kind in ("konu", "not")
        self.pending_page_break = False
        self.pending_full_page = False
        self.doc.blocks.append(block)
        self.block = block
        self.in_solution = False
        return block

    def ensure(self) -> Block:
        if self.block is None:
            self.start("soru")
        return self.block

    def target(self) -> list:
        block = self.ensure()
        return block.solution if self.in_solution else block.body

    # -- icerik ekleme --
    def add_para(self, runs: list[Run], kind: str = "paragraph") -> None:
        runs = _strip_runs(runs)
        if not runs:
            return
        if kind == "formula":
            # Formullerde yazarin koydugu coklu bosluk anlamlidir; bolunmez
            # bosluga cevirerek koruyoruz.
            runs = [
                Run(re.sub(r" {2,}", lambda m: " " * len(m.group(0)), r.text), r.bold, r.script)
                for r in runs
            ]
        self.target().append(Para(runs, kind))

    def add_image(self, path: str) -> None:
        self.target().append(Image(path))

    def add_option(self, letter: str, runs: list[Run]) -> None:
        block = self.ensure()
        block.options.append(Option(letter.upper(), _strip_runs(runs)))

    # -- tek paragrafin islenmesi --
    def feed(self, runs: list[Run], style_name: str = "", is_list: bool = False) -> None:
        runs = [Run(r.text, r.bold, r.script) for r in runs]
        text = "".join(r.text for r in runs).strip()

        # 1) kose parantezli isaretler
        while True:
            m = MARKER_RE.match(text)
            if not m:
                break
            name = m.group("name").strip().lower()
            text = text[m.end() :]
            runs = _consume_prefix(runs, m.end())
            if name in ("sayfa", "sayfa sonu", "yeni sayfa"):
                self.pending_page_break = True
            elif name in ("tam sayfa", "tamsayfa"):
                self.pending_full_page = True
            elif name in ("çözüm", "cozum"):
                self.ensure()
                self.in_solution = True
            elif name in ("formül", "formul"):
                self.add_para(runs, "formula")
                return
            elif name in ("başlık", "baslik", "ara başlık", "ara baslik"):
                self.add_para(runs, "subheading")
                return
            else:
                kind = _KIND_ALIASES.get(name)
                if kind:
                    self.start(kind)
                else:
                    self.doc.warnings.append(f"Bilinmeyen isaret: [{name}] (yok sayildi)")
            if not text.strip():
                return

        if not text:
            return

        # 2) Word baslik stilleri
        low_style = (style_name or "").lower()
        if low_style.startswith(("heading", "başlık", "baslik", "title")):
            level = _heading_level(low_style)
            if level <= 1:
                block = self.start("konu")
                block.title = _strip_runs(runs)
            else:
                self.add_para(runs, "subheading")
            return

        # 3) blok baslatan kaliplar
        m = _ORNEK_RE.match(text)
        if m and _looks_like_heading_prefix(text, m):
            number = int(m.group(2)) if m.group(2) else None
            self.start("ornek", number)
            runs = _consume_prefix(runs, m.end())
            text = text[m.end() :]
            if not text.strip():
                return

        elif (m := _SORU_RE.match(text)) and _looks_like_heading_prefix(text, m):
            number = int(m.group(1)) if m.group(1) else None
            self.start("soru", number)
            runs = _consume_prefix(runs, m.end())
            text = text[m.end() :]
            if not text.strip():
                return

        elif m := _KONU_RE.match(text):
            block = self.start("konu")
            runs = _consume_prefix(runs, m.end())
            text = text[m.end() :]
            if text.strip():
                block.title = _strip_runs(runs)
            return

        elif m := _NOT_RE.match(text):
            self.start("not")
            runs = _consume_prefix(runs, m.end())
            text = text[m.end() :]
            if not text.strip():
                return

        elif (m := _NUMBERED_RE.match(text)) and not is_list:
            # "1. ..." ile baslayan paragraf yeni bir SORU acar. (Word'un
            # otomatik numarali listeleri bu kuralin disindadir; onlar madde
            # olarak akar.) Acik ama bos bir blok varsa numara ona yazilir.
            acik = self.block
            if (
                acik is not None
                and acik.kind in ("soru", "ornek")
                and not (acik.body or acik.options or acik.solution)
            ):
                acik.number = int(m.group(1))
            else:
                self.start("soru", int(m.group(1)))
            runs = _consume_prefix(runs, m.end() - len(m.group(2)))
            text = m.group(2)
            if not text.strip():
                return

        # 4) cozum
        if _SOLUTION_RE.match(text) and self.block is not None:
            m = _SOLUTION_RE.match(text)
            self.in_solution = True
            runs = _consume_prefix(runs, m.end())
            text = text[m.end() :]
            if not text.strip():
                return

        # 5) secenek
        m = _OPTION_RE.match(text)
        if m and self.block is not None and self.block.kind in ("soru", "ornek"):
            offset = len(text) - len(m.group(2))
            self.add_option(m.group(1), _consume_prefix(runs, offset))
            return

        # 6) formul / madde / paragraf
        if _FORMULA_RE.match(text):
            # $$ ... $$ isaretlerini metinden degil Run'lardan kirp; boylece
            # icerideki ustsimge/altsimge bicimleri korunur.
            self.add_para(_trim_runs(_strip_runs(runs), 2, 2), "formula")
            return
        if is_list:
            self.add_para(runs, "bullet")
            return
        if m := _BULLET_RE.match(text):
            self.add_para(_consume_prefix(runs, len(text) - len(m.group(1))), "bullet")
            return
        self.add_para(runs, "paragraph")


def _heading_level(style_name: str) -> int:
    m = re.search(r"(\d+)", style_name)
    return int(m.group(1)) if m else 1


def _looks_like_heading_prefix(text: str, match: re.Match) -> bool:
    """'Örnek 3:' gibi bir blok basligini, 'Örnegin ...' gibi normal cumleden
    ayirir."""
    consumed = text[: match.end()]
    return bool(re.search(r"[:.\-–]\s*$", consumed) or match.group(0).strip().lower() in
                ("örnek", "ornek", "soru")) or bool(
        re.match(r"^(örnek|ornek|soru)\s*\d+", consumed, re.IGNORECASE)
    )


def _trim_runs(runs: list[Run], head: int, tail: int) -> list[Run]:
    """Run listesinin basindan `head`, sonundan `tail` karakter kirpar."""
    out = _consume_prefix(runs, head)
    remaining = tail
    for i in range(len(out) - 1, -1, -1):
        if remaining <= 0:
            break
        text = out[i].text
        cut = min(len(text), remaining)
        remaining -= cut
        out[i] = Run(text[: len(text) - cut], out[i].bold, out[i].script)
    return [r for r in out if r.text]


def _consume_prefix(runs: list[Run], count: int) -> list[Run]:
    """Metnin basindan `count` karakter kirpar, bicimleri korur."""
    out: list[Run] = []
    remaining = count
    for run in runs:
        if remaining <= 0:
            out.append(run)
            continue
        if len(run.text) <= remaining:
            remaining -= len(run.text)
            continue
        out.append(Run(run.text[remaining:], run.bold, run.script))
        remaining = 0
    return out


# --------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------
def parse_docx(
    path: str | Path,
    media_dir: Path | None = None,
    continuous_numbering: bool = False,
) -> Document:
    from docx import Document as DocxDocument
    from docx.oxml.ns import qn

    document = DocxDocument(str(path))
    media_dir = Path(media_dir or tempfile.mkdtemp(prefix="mathera-media-"))
    media_dir.mkdir(parents=True, exist_ok=True)

    builder = _Builder()
    part = document.part
    body = document.element.body

    from docx.text.paragraph import Paragraph
    from docx.table import Table

    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, document)
            _feed_paragraph(builder, para, part, media_dir, qn)
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        _feed_paragraph(builder, para, part, media_dir, qn)
            builder.doc.warnings.append(
                "Word tablosu bulundu: hucre icerikleri sirayla duz paragraf olarak alindi."
            )
    return finalize(builder.doc, continuous_numbering)


def _feed_paragraph(builder: _Builder, para, part, media_dir: Path, qn) -> None:
    # gorseller
    for run in para.runs:
        for blip in run._element.findall(".//" + qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            if not rid:
                continue
            image_part = part.related_parts[rid]
            target = media_dir / Path(image_part.partname).name
            target.write_bytes(image_part.blob)
            builder.add_image(str(target))

    runs: list[Run] = []
    for run in para.runs:
        if not run.text:
            continue
        script = ""
        font = run.font
        if font.superscript:
            script = "sup"
        elif font.subscript:
            script = "sub"
        runs.extend(parse_inline(run.text, bold=bool(run.bold), script=script))

    # Word denklem nesneleri (OMML): duz metnini al
    for node in para._element.findall(".//" + qn("m:t")):
        if node.text:
            runs.append(Run(node.text))
            builder.doc.warnings.append(
                "Word denklem nesnesi duz metne cevrildi; ustsimge/altsimge icin "
                "^{} ve _{} yazimini kullanin."
            )

    if not runs:
        return
    style_name = para.style.name if para.style is not None else ""
    is_list = _is_list_paragraph(para, style_name, qn)
    builder.feed(runs, style_name, is_list)


def _is_list_paragraph(para, style_name: str, qn) -> bool:
    if para._element.find(".//" + qn("w:numPr")) is not None:
        return True
    low = (style_name or "").lower()
    return "list" in low or "madde" in low


# --------------------------------------------------------------------------
# Duz metin / markdown
# --------------------------------------------------------------------------
def parse_text(text: str, continuous_numbering: bool = False) -> Document:
    builder = _Builder()
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        style = ""
        is_list = False
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            style = f"heading {level}"
            line = line.lstrip("#").strip()
        elif _BULLET_RE.match(line):
            is_list = True
            line = _BULLET_RE.match(line).group(1)
        builder.feed(parse_inline(line), style, is_list)
    return finalize(builder.doc, continuous_numbering)


# --------------------------------------------------------------------------
# JSON (LLM'in urettigi kanonik bicim)
# --------------------------------------------------------------------------
def parse_json(path: str | Path, continuous_numbering: bool = False) -> Document:
    with open(path, encoding="utf-8") as fh:
        data = json.load(fh)
    doc = Document()
    for item in data.get("bloklar", data.get("blocks", [])):
        kind = _KIND_ALIASES.get(str(item.get("tur", item.get("kind", "soru"))).lower(), "soru")
        block = Block(
            kind=kind,
            number=item.get("no", item.get("number")),
            force_full_page=bool(item.get("tam_sayfa", False)) or kind in ("konu", "not"),
            page_break_before=bool(item.get("sayfa_basi", False)),
        )
        if item.get("baslik") or item.get("title"):
            block.title = parse_inline(item.get("baslik") or item["title"])
        for entry in item.get("govde", item.get("body", [])):
            _json_flowable(block.body, entry)
        for entry in item.get("cozum", item.get("solution", [])):
            _json_flowable(block.solution, entry)
        for opt in item.get("secenekler", item.get("options", [])):
            if isinstance(opt, str):
                m = _OPTION_RE.match(opt)
                if m:
                    block.options.append(Option(m.group(1).upper(), parse_inline(m.group(2))))
            else:
                block.options.append(
                    Option(
                        str(opt.get("harf", opt.get("letter", "A"))).upper(),
                        parse_inline(str(opt.get("metin", opt.get("text", "")))),
                    )
                )
        doc.blocks.append(block)
    return finalize(doc, continuous_numbering)


def _json_flowable(target: list, entry) -> None:
    if isinstance(entry, str):
        target.append(Para(parse_inline(entry), "paragraph"))
        return
    kind = entry.get("tur", entry.get("kind", "paragraph"))
    kind = {
        "paragraf": "paragraph",
        "madde": "bullet",
        "formul": "formula",
        "formül": "formula",
        "ara_baslik": "subheading",
        "gorsel": "image",
        "görsel": "image",
    }.get(kind, kind)
    if kind == "image":
        target.append(Image(entry["yol"] if "yol" in entry else entry["path"]))
        return
    target.append(Para(parse_inline(str(entry.get("metin", entry.get("text", "")))), kind))


# --------------------------------------------------------------------------
def finalize(doc: Document, continuous_numbering: bool = False) -> Document:
    """Numaralari tamamlar, bos bloklari eler.

    Varsayilan: soru ve ornek ayri ayri numaralanir (1., 2., ... her tur icin).
    ``continuous_numbering`` ile ikisi tek bir dizide numaralanir.
    """
    doc.blocks = [
        b for b in doc.blocks if b.body or b.options or b.solution or b.title
    ]
    counters = {"soru": 0, "ornek": 0}
    for block in doc.blocks:
        if block.kind not in counters:
            continue
        key = "soru" if continuous_numbering else block.kind
        if block.number is None:
            counters[key] += 1
            block.number = counters[key]
        else:
            counters[key] = block.number
        if continuous_numbering:
            counters["ornek"] = counters["soru"]
    return doc


def parse(path: str | Path, continuous_numbering: bool = False) -> Document:
    """Uzantiya gore dogru cozumleyiciyi secer."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return parse_docx(path, continuous_numbering=continuous_numbering)
    if suffix == ".json":
        return parse_json(path, continuous_numbering)
    if suffix in (".txt", ".md", ".markdown"):
        return parse_text(path.read_text(encoding="utf-8"), continuous_numbering)
    raise ValueError(f"Desteklenmeyen girdi bicimi: {suffix}")
