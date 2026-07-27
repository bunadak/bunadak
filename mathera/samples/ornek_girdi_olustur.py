"""Ornek Word girdisi uretir: `python3 samples/ornek_girdi_olustur.py`

Uretilen dosya (samples/ornek_girdi.docx) motorun tanidigi tum yapilari
gosterir: konu anlatimi, not, ornek + cozum, soru + secenekler, madde
listesi, formul, ustsimge/altsimge.
"""

from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent


def build(path: Path) -> Path:
    doc = Document()

    doc.add_paragraph("[KONU]")
    doc.add_paragraph("İkinci Dereceden Denklemler", style="Heading 1")
    doc.add_paragraph(
        "a, b ve c gerçek sayılar ve a ≠ 0 olmak üzere ax^2 + bx + c = 0 biçiminde "
        "yazılabilen denklemlere ikinci dereceden bir bilinmeyenli denklem denir. "
        "Bu denklemin kökleri, denklemi sağlayan x değerleridir."
    )
    doc.add_paragraph("Diskriminant", style="Heading 2")
    doc.add_paragraph("$$Δ = b^2 − 4ac$$")
    doc.add_paragraph("Δ > 0 ise denklemin birbirinden farklı iki gerçek kökü vardır.", style="List Bullet")
    doc.add_paragraph("Δ = 0 ise denklemin eşit iki (çakışık) kökü vardır.", style="List Bullet")
    doc.add_paragraph("Δ < 0 ise denklemin gerçek kökü yoktur.", style="List Bullet")
    doc.add_paragraph(
        "Kökler toplamı ve kökler çarpımı, katsayılar cinsinden doğrudan yazılabilir; "
        "bu bağıntılar kökleri bulmadan da soru çözmeyi mümkün kılar."
    )
    doc.add_paragraph("$$x_1 + x_2 = −b/a        x_1 · x_2 = c/a$$")

    doc.add_paragraph("[NOT]")
    doc.add_paragraph(
        "Δ = 0 durumunda kök **çift katlı** yazılır: x_1 = x_2 = −b/(2a). "
        "Bir denklemin eşit iki kökü olduğu söyleniyorsa, çözüme her zaman "
        "Δ = 0 eşitliğinden başlanır."
    )

    doc.add_paragraph("[ÖRNEK]")
    doc.add_paragraph(
        "x^2 − 6x + m − 2 = 0 denkleminin eşit iki gerçek kökü olduğuna göre, "
        "m kaçtır?"
    )
    doc.add_paragraph("Çözüm:")
    doc.add_paragraph("Eşit iki kök koşulu Δ = 0 demektir.")
    doc.add_paragraph("$$Δ = (−6)^2 − 4 · 1 · (m − 2) = 0$$")
    doc.add_paragraph("36 − 4m + 8 = 0 ⇒ 4m = 44 ⇒ m = 11 bulunur.")

    doc.add_paragraph("[ÖRNEK]")
    doc.add_paragraph(
        "Kökleri 3 ve −5 olan ikinci dereceden denklemi yazınız."
    )
    doc.add_paragraph("Çözüm:")
    doc.add_paragraph(
        "Kökler toplamı 3 + (−5) = −2, kökler çarpımı 3 · (−5) = −15 olur."
    )
    doc.add_paragraph("$$x^2 + 2x − 15 = 0$$")

    doc.add_paragraph("[SORU]")
    doc.add_paragraph(
        "x^2 − 6x + m − 2 = 0 denkleminin eşit iki gerçek kökü olduğuna göre, m kaçtır?"
    )
    for letter, value in zip("ABCDE", ["7", "9", "11", "13", "15"]):
        doc.add_paragraph(f"{letter}) {value}")

    doc.add_paragraph("[SORU]")
    doc.add_paragraph(
        "Bir manav elindeki elmaların önce %20'sini, sonra kalanın %25'ini satıyor. "
        "Geriye 240 elma kaldığına göre, manavın başlangıçtaki elma sayısı kaçtır?"
    )
    for letter, value in zip("ABCDE", ["360", "380", "400", "420", "440"]):
        doc.add_paragraph(f"{letter}) {value}")

    doc.add_paragraph("[SORU]")
    doc.add_paragraph(
        "A = {1, 2, 3, 4, 5} kümesinin alt kümelerinin kaç tanesinde 2 bulunur, "
        "3 bulunmaz?"
    )
    for letter, value in zip("ABCDE", ["4", "6", "8", "12", "16"]):
        doc.add_paragraph(f"{letter}) {value}")

    doc.save(str(path))
    return path


if __name__ == "__main__":
    print(build(HERE / "ornek_girdi.docx"))
